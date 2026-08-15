#!/usr/bin/env python3
"""Patch one copy and rerun it with a static --output; the Host observes it automatically."""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    parser.add_argument("--title", default="Document title")
    parser.add_argument(
        "--image",
        action="append",
        default=[],
        help="Input mountPath relative to MYCOPILOT_INPUT_ROOT; may be repeated.",
    )
    return parser.parse_args()


def mounted_input(mount_path: str) -> Path:
    root_value = os.environ.get("MYCOPILOT_INPUT_ROOT")
    if not root_value:
        raise RuntimeError("MYCOPILOT_INPUT_ROOT is required when --image is used")
    root = Path(root_value).resolve()
    relative = Path(mount_path)
    if relative.is_absolute():
        raise ValueError(f"input mountPath must be relative: {mount_path}")
    candidate = (root / relative).resolve()
    if candidate != root and root not in candidate.parents:
        raise ValueError(f"input escapes MYCOPILOT_INPUT_ROOT: {mount_path}")
    if not candidate.is_file():
        raise FileNotFoundError(candidate)
    return candidate


def set_font(run, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def build_document(title: str, images: list[Path]) -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    heading = document.add_paragraph()
    heading.style = document.styles["Title"]
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(heading.add_run(title), "Microsoft YaHei", 24, bold=True)

    # Patch this content block; keep one builder and rerun it after every revision.
    body = document.add_paragraph()
    body.paragraph_format.space_after = Pt(10)
    set_font(
        body.add_run("Replace this paragraph with the requested document content."),
        "Microsoft YaHei",
        10.5,
    )

    for image_path in images:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Inches(6.1))
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(caption.add_run(image_path.stem), "Microsoft YaHei", 9)

    return document


def publish(document: Document, output: Path) -> None:
    if output.suffix.lower() != ".docx":
        raise ValueError("--output must end in .docx")
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.name}.", suffix=".docx", dir=output.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        document.save(temporary)
        Document(temporary)  # Re-open before publication to catch a malformed package.
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def main() -> None:
    args = arguments()
    output = Path(args.output).resolve()
    images = [mounted_input(value) for value in args.image]
    publish(build_document(args.title, images), output)
    print(
        json.dumps(
            {"status": "created", "output": str(output), "images": len(images)},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
