#!/usr/bin/env python3
"""Edit one mounted DOCX and save a distinct output through the managed Host."""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from pathlib import Path

from docx import Document


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        required=True,
        help="Logical input mountPath below MYCOPILOT_INPUT_ROOT.",
    )
    parser.add_argument("--output", required=True, help="Distinct save-as .docx output.")
    return parser.parse_args()


def mounted_input(mount_path: str) -> Path:
    root_value = os.environ.get("MYCOPILOT_INPUT_ROOT")
    if not root_value:
        raise RuntimeError("MYCOPILOT_INPUT_ROOT is required")
    relative = Path(mount_path)
    if relative.is_absolute():
        raise ValueError(f"input mountPath must be relative: {mount_path}")
    root = Path(root_value).resolve(strict=True)
    candidate = (root / relative).resolve(strict=True)
    if candidate == root or root not in candidate.parents:
        raise ValueError(f"input escapes MYCOPILOT_INPUT_ROOT: {mount_path}")
    if not candidate.is_file():
        raise FileNotFoundError(candidate)
    return candidate


def edit_document(document, source_path: Path) -> None:
    """Patch only the marked region; normal Python and pinned imports are available."""

    # BEGIN EDIT REGION
    # Example:
    # for paragraph in document.paragraphs:
    #     if "Old text" in paragraph.text:
    #         for run in paragraph.runs:
    #             run.text = run.text.replace("Old text", "New text")
    #
    # Fixed managed dependencies may be imported here. Nested helper functions,
    # loops, conditions, and data transformations are ordinary Python.
    raise RuntimeError("Replace the EDIT REGION with the requested document edits")
    # END EDIT REGION


def publish(document, source: Path, output: Path) -> None:
    if output.suffix.lower() != ".docx":
        raise ValueError("--output must end in .docx")
    output = output.resolve()
    if output == source:
        raise ValueError("--output must be distinct from the mounted --source")
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.name}.", suffix=".docx", dir=output.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        document.save(temporary)
        Document(temporary)  # Reopen the candidate before publication.
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def main() -> None:
    args = arguments()
    source = mounted_input(args.source)
    if source.suffix.lower() != ".docx":
        raise ValueError("--source must name one mounted .docx file")
    output = Path(args.output)
    document = Document(source)
    edit_document(document, source)
    publish(document, source, output)
    print(
        json.dumps(
            {"status": "edited", "source": args.source, "output": str(output)},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
