#!/usr/bin/env python3
"""Edit one frozen XLSX input and save a distinct managed candidate output.

Patch only ``edit_workbook``. This is ordinary Python, not a declarative plan:
helpers, loops, conditions, comprehensions, and imports from the pinned runtime
are allowed. The Host freezes this file and the inputs before execution, rewrites
the output to private staging, validates the result, then publishes atomically.
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path

from openpyxl import load_workbook


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        required=True,
        help="Input mountPath relative to MYCOPILOT_INPUT_ROOT.",
    )
    parser.add_argument("--output", required=True)
    return parser.parse_args()


def mounted_input(mount_path: str) -> tuple[Path, Path]:
    root_value = os.environ.get("MYCOPILOT_INPUT_ROOT")
    if not root_value:
        raise RuntimeError("MYCOPILOT_INPUT_ROOT is required")
    root = Path(root_value).resolve()
    relative = Path(mount_path)
    if relative.is_absolute():
        raise ValueError(f"input mountPath must be relative: {mount_path}")
    source = (root / relative).resolve()
    if source != root and root not in source.parents:
        raise ValueError(f"input escapes MYCOPILOT_INPUT_ROOT: {mount_path}")
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        raise ValueError("--source must identify one mounted .xlsx file")
    return root, source


def edit_workbook(workbook, input_root: Path) -> None:
    """Apply the requested edit using normal Python and the pinned libraries."""

    # BEGIN EDIT REGION
    # Replace this line with task-specific Python. Example:
    # sheet = workbook["预算"]
    # for row in range(2, sheet.max_row + 1):
    #     sheet.cell(row=row, column=5, value=f"=SUM(B{row}:D{row})")
    raise RuntimeError("Replace the EDIT REGION with the requested workbook edits")
    # END EDIT REGION


def publish(workbook, output: Path) -> None:
    if output.suffix.lower() != ".xlsx":
        raise ValueError("--output must end in .xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)
    # The Host rewrites this path to its private candidate and publishes atomically.
    workbook.save(output)


def main() -> None:
    args = arguments()
    input_root, source = mounted_input(args.source)
    output = Path(args.output).resolve()
    workbook = load_workbook(
        source,
        data_only=False,
        read_only=False,
        keep_links=True,
        rich_text=True,
    )
    try:
        edit_workbook(workbook, input_root)
        publish(workbook, output)
    finally:
        workbook.close()
    print(json.dumps({"status": "edited", "output": str(output)}))


if __name__ == "__main__":
    main()
