#!/usr/bin/env python3
"""Patch one copy and rerun it with a static --output; the Host observes it automatically."""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill


def mounted_input(mount_path: str) -> Path:
    """Resolve one declared run_command.inputs mount for task-specific Builder code."""

    root_value = os.environ.get("MYCOPILOT_INPUT_ROOT")
    if not root_value:
        raise RuntimeError("MYCOPILOT_INPUT_ROOT is required for mounted inputs")
    root = Path(root_value).resolve()
    relative = Path(mount_path)
    if relative.is_absolute():
        raise ValueError(f"input mountPath must be relative: {mount_path}")
    candidate = (root / relative).resolve(strict=True)
    if candidate != root and root not in candidate.parents:
        raise ValueError(f"input escapes MYCOPILOT_INPUT_ROOT: {mount_path}")
    if not candidate.is_file():
        raise FileNotFoundError(candidate)
    return candidate


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    return parser.parse_args()


def build_workbook():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "数据"

    # Patch this content block; keep formulas as formulas and values as typed values.
    rows = [
        ["项目", "1月", "2月", "3月", "合计"],
        ["A", 1200, 1300, 1250, None],
        ["B", 900, 1100, 1050, None],
        ["C", 700, 800, 950, None],
    ]
    for row in rows:
        worksheet.append(row)
    for row_number in range(2, len(rows) + 1):
        worksheet.cell(row=row_number, column=5).value = (
            f"=SUM(B{row_number}:D{row_number})"
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for column in range(2, 6):
        for row_number in range(2, len(rows) + 1):
            worksheet.cell(row=row_number, column=column).number_format = "¥#,##0.00"
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 18
    for name in ("B", "C", "D", "E"):
        worksheet.column_dimensions[name].width = 14
    worksheet.conditional_formatting.add(
        f"E2:E{len(rows)}",
        ColorScaleRule(
            start_type="min",
            start_color="FEE2E2",
            mid_type="percentile",
            mid_value=50,
            mid_color="FEF3C7",
            end_type="max",
            end_color="DCFCE7",
        ),
    )

    chart = BarChart()
    chart.type = "col"
    chart.title = "合计"
    chart.add_data(
        Reference(worksheet, min_col=5, min_row=1, max_row=len(rows)),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(worksheet, min_col=1, min_row=2, max_row=len(rows))
    )
    worksheet.add_chart(chart, "G2")
    return workbook


def publish(workbook, output: Path) -> None:
    if output.suffix.lower() != ".xlsx":
        raise ValueError("--output must end in .xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.name}.", suffix=".xlsx", dir=output.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        verified = load_workbook(temporary, data_only=False, read_only=False)
        if not verified.sheetnames:
            raise RuntimeError("published workbook has no worksheets")
        verified.close()
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def main() -> None:
    args = arguments()
    output = Path(args.output).resolve()
    workbook = build_workbook()
    publish(workbook, output)
    print(
        json.dumps(
            {"status": "created", "output": str(output)},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
